import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .models import Production, Task, Subtask, Expense, TeamMember, AuditLog

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")


def _add_sheet(wb, title, columns, rows):
    ws = wb.create_sheet(title=title)
    ws.append(columns)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for row in rows:
        ws.append(row)
    for i, col in enumerate(columns, start=1):
        width = max(len(str(col)), *(len(str(v)) for v in [r[i - 1] for r in rows] + [""]))
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(width + 2, 40)
    ws.freeze_panes = "A2"


def build_export_workbook():
    """Build an in-memory .xlsx workbook of the app's data, one sheet per
    entity. Login credentials (password hashes) and avatar images are left
    out — this is a data export, not a account/credentials dump."""
    wb = Workbook()
    wb.remove(wb.active)

    _add_sheet(
        wb, "Productions",
        ["ID", "Client", "Shoot Name", "Type", "Status", "Shoot Date", "Start Time", "End Time", "Budget",
         "Production Hours", "Looks/SKUs", "Frame Count", "Videos", "Location",
         "Crew", "Drive Link", "Brief Link", "Rundown Link", "Notes", "Created At", "Updated At"],
        [[p.id, p.client, p.shoot_name, p.type, p.status,
          p.shoot_date.isoformat() if p.shoot_date else "", p.shoot_time_start or "", p.shoot_time_end or "",
          p.budget, p.production_hours,
          p.looks_skus, p.frame_count, p.videos, p.location,
          ", ".join(c.name for c in p.crew), p.drive_link, p.brief_link, p.rundown_link,
          p.notes, p.created_at.isoformat(), p.updated_at.isoformat()]
         for p in Production.query.order_by(Production.id).all()],
    )

    _add_sheet(
        wb, "Tasks",
        ["ID", "Title", "Production", "Assignees", "Created By", "Due Date",
         "Status", "Priority", "Notes", "Created At", "Updated At"],
        [[t.id, t.title,
          f"{t.production.client} — {t.production.shoot_name}" if t.production else "",
          ", ".join(m.name for m in t.assignees), t.created_by.name if t.created_by else "",
          t.due_date.isoformat() if t.due_date else "", t.status, t.priority, t.notes,
          t.created_at.isoformat(), t.updated_at.isoformat()]
         for t in Task.query.order_by(Task.id).all()],
    )

    _add_sheet(
        wb, "Subtasks",
        ["ID", "Task ID", "Task Title", "Text", "Done", "Position"],
        [[s.id, s.task_id, s.task.title if s.task else "", s.text, "Yes" if s.done else "No", s.position]
         for s in Subtask.query.order_by(Subtask.task_id, Subtask.position).all()],
    )

    _add_sheet(
        wb, "Expenses",
        ["ID", "Production", "Description", "Category", "Brand", "Amount", "Date",
         "Paid By", "Status", "Notes", "Created At"],
        [[e.id, f"{e.production.client} — {e.production.shoot_name}" if e.production else "",
          e.description, e.category, e.brand or "", e.amount, e.date.isoformat() if e.date else "",
          e.paid_by.name if e.paid_by else "", e.status, e.notes, e.created_at.isoformat()]
         for e in Expense.query.order_by(Expense.id).all()],
    )

    _add_sheet(
        wb, "Team Members",
        ["ID", "Name", "Username", "Role", "Has Login", "Last Signed In", "Created At"],
        [[m.id, m.name, m.username or "", m.role, "Yes" if m.username else "No",
          m.last_login_at.isoformat() if m.last_login_at else "", m.created_at.isoformat()]
         for m in TeamMember.query.order_by(TeamMember.id).all()],
    )

    _add_sheet(
        wb, "Changes Log",
        ["ID", "When", "Who", "Action", "Entity Type", "Entity", "Detail"],
        [[a.id, a.created_at.isoformat(), a.actor_name, a.action, a.entity_type, a.entity_label, a.detail or ""]
         for a in AuditLog.query.order_by(AuditLog.created_at.desc()).limit(1000).all()],
    )

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
